# Copyright (c) 2026 Ashutosh Sinha <ajsinha@gmail.com>. All rights reserved.
# Proprietary and confidential; see LICENSE at the repository root.
"""
Slide layouts, each drawn from a plain dictionary.

A deck is data: a list of slide specs, each naming its ``kind``. Keeping the
content out of the drawing code is what lets one deck be assembled from several
modules that share one set of layouts, and what keeps every layout short enough
to read -- and to fit.

Kinds: ``title``, ``act`` (and the older ``divider``), ``qa``, ``steps``, ``bullets``,
``table``, ``cards``, ``stats``, ``split``, ``flow``, ``context``, ``code``, ``tree``, ``shot``,
``shots``, ``thanks``.

Every spec may carry ``source``: it begins the slide's speaker notes, and names
the file each figure and claim on the slide comes from; ``talk`` follows it as the talk
track. ``takeaway`` puts the slide's one-sentence key insight in a band above the footer.

Project Pravaha -- Ask once. Answer always.
Copyright (c) 2026 Ashutosh Sinha <ajsinha@gmail.com>. All rights reserved.
PROPRIETARY AND CONFIDENTIAL. See LICENSE at the repository root.
"""

from __future__ import annotations

from typing import Any

from pptx.enum.text import PP_ALIGN

import theme as T
from metrics import HEAD, SANS, SH, SW, text_h

GAP = 0.18


def _items(tf: Any, items: list[Any], size: float) -> None:
    """Bulleted items; a tuple is ``(head, body)``."""
    for i, it in enumerate(items):
        first = i == 0
        if isinstance(it, tuple):
            T.runs(
                tf,
                [("▪  ", T.ACCENT, True), (it[0], T.INK, True)],
                size=size,
                space_after=1,
                first=first,
                space_before=0 if first else 5,
            )
            T.runs(tf, [(it[1], T.SLATE, False)], size=size - 1.5, space_after=0, level=1)
        else:
            T.runs(
                tf,
                [("▪  ", T.ACCENT, True), (it, T.INK, False)],
                size=size,
                space_after=0,
                first=first,
                space_before=0 if first else 6,
            )


def _intro(sl: Any, y: float, text: str | None) -> float:
    if not text:
        return y
    used = T.fitted(
        sl,
        T.ML,
        y,
        T.CW,
        1.05,
        lambda tf, s: T.para(tf, text, size=s, color=T.INK, first=True, space_after=0, line=1.25),
        14,
        10.5,
    )
    return y + used + GAP


def _takeaway(sl: Any, text: str) -> float:
    """The slide's one-sentence takeaway, in a tinted band pinned above the footer; returns
    the top of the space left above it."""
    label = "Key insight  "
    width = T.CW - 0.5
    size = 14.0
    while size > 10 and text_h(label + text, width * 0.94, size, True, SANS, 1.2) > 0.75:
        size -= 0.5
    h = text_h(label + text, width * 0.94, size, True, SANS, 1.2) + 0.24
    top = T.BODY_BOTTOM - h
    T.rect(sl, T.ML, top, T.CW, h, fill=T.WASH)
    T.rect(sl, T.ML, top, 0.06, h, fill=T.ACCENT)
    T.fitted(
        sl,
        T.ML + 0.28,
        top + 0.12,
        width,
        h - 0.2,
        lambda tf, s: T.runs(tf, [(label, T.ACCENT, True), (text, T.ACCENT_DD, True)], size=s,
                             first=True, space_after=0, line=1.2),
        size,
        9.5,
    )
    return top - GAP


def _note(sl: Any, text: str | None, takeaway: str | None = None) -> float:
    """A crimson-ruled note pinned above the footer -- or, for a ``takeaway``, the key-insight
    band; returns the top of the space left above it."""
    if takeaway:
        return _takeaway(sl, takeaway)
    if not text:
        return T.BODY_BOTTOM
    size = 12.0
    width = T.CW - 0.3
    while size > 9 and text_h(text, width * 0.94, size, False, SANS, 1.2) > 0.9:
        size -= 0.5
    h = text_h(text, width * 0.94, size, False, SANS, 1.2)
    top = T.BODY_BOTTOM - h
    T.rect(sl, T.ML, top, 0.045, h, fill=T.ACCENT)
    T.fitted(
        sl,
        T.ML + 0.25,
        top,
        width,
        h + 0.02,
        lambda tf, s: T.para(
            tf, text, size=s, color=T.SLATE, italic=True, first=True, space_after=0, line=1.2
        ),
        size,
        8.5,
    )
    return top - GAP


def _body(sl: Any, x: float, y: float, w: float, h: float, spec: dict[str, Any],
          start: float = 15.0) -> None:
    """Items or a table in a box -- whichever the spec carries."""
    if spec.get("rows"):
        T.fitted_table(sl, spec["rows"], x, y, w, h, spec.get("col_w"),
                       start=spec.get("size", 13), bold_col0=spec.get("bold_col0", True))
    elif spec.get("items"):
        T.fitted(sl, x, y, w, h, lambda tf, size: _items(tf, spec["items"], size),
                 spec.get("size", start), 9)


def title(s: dict[str, Any]) -> None:
    T._state["n"] = 1
    sl = T.blank()
    T.rect(sl, 0, 0, SW, SH, fill=T.WHITE)
    T.rect(sl, 0, 0, SW, 4.3, fill=T.ACCENT_D)
    T.rect(sl, 0, 4.3, SW, 0.06, fill=T.ACCENT_L)
    T.rect(sl, 0, 0, 0.20, 4.3, fill=T.ACCENT_DD)
    T.mark(sl, SW - T.ML - 1.55, 0.75, 1.5, light=True)
    tf = T.txt(sl, T.ML + 0.3, 0.8, T.CW * 0.7, 0.34)
    T.para(tf, s["kicker"], size=11, color=T.ACCENT_L, bold=True, first=True, space_after=0)

    def head(tf: Any, size: float) -> None:
        for i, line in enumerate(s["title"]):
            T.para(tf, line, size=size, color=T.WHITE, font=HEAD, bold=True, first=i == 0,
                   space_after=0, line=1.08)

    T.fitted(sl, T.ML + 0.3, 1.3, T.CW * 0.78, 1.85, head, 38, 24)
    T.rect(sl, T.ML + 0.3, 3.25, 1.7, 0.035, fill=T.ACCENT_L)
    # The slogan, set to be read rather than noticed in passing (brand/README.md: it belongs
    # on the logo lockup, the landing page, the README and the sign-in screen -- and a title).
    tf = T.txt(sl, T.ML + 0.3, 3.42, T.CW * 0.85, 0.7)
    T.para(tf, s["sub"], size=25, color=T.WHITE, italic=True, font=HEAD, first=True,
           space_after=0)
    tf = T.txt(sl, T.ML + 0.3, 4.75, T.CW * 0.4, 1.3)
    T.para(tf, "Ashutosh Sinha", size=20, color=T.INK, bold=True, font=HEAD, first=True,
           space_after=3)
    T.para(tf, s.get("date", "September 2026"), size=12, color=T.ACCENT, space_after=1)
    T.para(tf, s["version"], size=11, color=T.SLATE, space_after=0)
    x0 = T.ML + T.CW * 0.45
    tf = T.txt(sl, x0, 4.62, T.CW * 0.5, 0.24)
    T.para(tf, "IN THIS DECK", size=9.5, color=T.ACCENT, bold=True, first=True, space_after=0)
    items = list(enumerate(s["agenda"], 1))
    half = (len(items) + 1) // 2
    colw = T.CW * 0.55 / 2 - 0.1
    for k, chunk in enumerate((items[:half], items[half:])):

        def agenda(tf: Any, size: float, chunk: list = chunk) -> None:
            for j, (i, c) in enumerate(chunk):
                T.runs(tf, [(f"{i}  ", T.ACCENT, True), (c, T.SLATE, False)], size=size,
                       space_after=3, first=j == 0)

        T.fitted(sl, x0 + k * (colw + 0.2), 4.92, colw, 1.9, agenda, 11, 8.5)
    T.footer(sl)
    T.notes(sl, s.get("source"), s.get("talk"))


def divider(s: dict[str, Any]) -> None:
    sl = T.divider(s["num"], s["title"], s["sub"], s["points"])
    T.notes(sl, s.get("source"), s.get("talk"))


def bullets(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    T.fitted(sl, T.ML, y, T.CW, bottom - y, lambda tf, size: _items(tf, s["items"], size),
             s.get("size", 17), 9.5)
    T.notes(sl, s.get("source"), s.get("talk"))


def table(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    T.fitted_table(sl, s["rows"], T.ML, y, T.CW, bottom - y, s.get("col_w"),
                   start=s.get("size", 14), bold_col0=s.get("bold_col0", True))
    T.notes(sl, s.get("source"), s.get("talk"))


def cards(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    items = s["cards"]
    cols = s.get("cols", 3)
    rows = (len(items) + cols - 1) // cols
    cw = (T.CW - GAP * (cols - 1)) / cols
    ch = (bottom - y - GAP * (rows - 1)) / rows
    for i, (num, head, body) in enumerate(items):
        r, c = divmod(i, cols)
        T.card(sl, T.ML + c * (cw + GAP), y + r * (ch + GAP), cw, ch, num, head, body)
    T.notes(sl, s.get("source"), s.get("talk"))


def stats(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    T.statbar(sl, y, s["stats"])
    y += 1.25 + GAP + 0.05
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    _body(sl, T.ML, y, T.CW, bottom - y, s, 16)
    T.notes(sl, s.get("source"), s.get("talk"))


def _column(sl: Any, x: float, y: float, w: float, h: float, col: dict[str, Any]) -> None:
    T.rect(sl, x, y, w, 0.42, fill=T.WASH_L)
    T.rect(sl, x, y, 0.045, 0.42, fill=T.ACCENT)
    tf = T.txt(sl, x + 0.18, y + 0.08, w - 0.3, 0.3)
    T.para(tf, col["head"], size=13, color=T.ACCENT_D, bold=True, first=True, space_after=0)
    top = y + 0.42 + 0.14
    _body(sl, x, top, w, y + h - top, col, 15)


def split(s: dict[str, Any]) -> None:
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    w = (T.CW - 0.4) / 2
    _column(sl, T.ML, y, w, bottom - y, s["left"])
    _column(sl, T.ML + w + 0.4, y, w, bottom - y, s["right"])
    T.notes(sl, s.get("source"), s.get("talk"))


def flow(s: dict[str, Any]) -> None:
    """Boxes in a row, joined by straight arrows; then optional items or a table beneath."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    steps = s["steps"]
    arrow = 0.30
    bw = (T.CW - arrow * (len(steps) - 1)) / len(steps)
    bh = s.get("box_h", 1.9)
    for i, (head, body) in enumerate(steps):
        x = T.ML + i * (bw + arrow)
        T.card(sl, x, y, bw, bh, "", head, body)
        if i:
            T.connect(sl, x - arrow + 0.04, y + bh / 2, x - 0.04, y + bh / 2, T.ACCENT, 2.0)
    top = y + bh + GAP + 0.05
    _body(sl, T.ML, top, T.CW, bottom - top, s, 15)
    T.notes(sl, s.get("source"), s.get("talk"))


def code(s: dict[str, Any]) -> None:
    """A code panel on the left and what it means on the right (or beneath, when
    ``wide`` is set)."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    if s.get("wide"):
        used = T.code(sl, T.ML, y, T.CW, s.get("code_h", (bottom - y) * 0.55), s["code"],
                      s.get("code_size", 13))
        top = y + used + GAP + 0.05
        _body(sl, T.ML, top, T.CW, bottom - top, s, 15)
    else:
        frac = s.get("code_w", 0.56)
        cw = T.CW * frac - 0.2
        T.code(sl, T.ML, y, cw, bottom - y, s["code"], s.get("code_size", 13))
        x = T.ML + T.CW * frac + 0.2
        _body(sl, x, y, T.ML + T.CW - x, bottom - y, s, 15)
    T.notes(sl, s.get("source"), s.get("talk"))


def context(s: dict[str, Any]) -> None:
    """A diagram: boxes placed on the content area, joined by arrows.

    Positions are fractions of the content box rather than inches, so a diagram keeps its
    proportions if the theme's margins move. An arrow leaves the side that faces its
    target: if the two boxes share a column the arrow is vertical, down the middle of the
    overlap; if they share a row it is horizontal; otherwise it runs centre to centre.
    """
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    x0, w0, h0 = T.ML, T.CW, bottom - y

    box: dict[str, tuple[float, float, float, float]] = {}
    for n in s["nodes"]:
        box[n["id"]] = (x0 + n["x"] * w0, y + n["y"] * h0, n["w"] * w0, n["h"] * h0)

    # Arrows first, so that a box always sits on top of the line that reaches it.
    for e in s.get("edges", []):
        ax, ay, aw, ah = box[e[0]]
        bx, by, bw, bh = box[e[1]]
        acx, acy, bcx, bcy = ax + aw / 2, ay + ah / 2, bx + bw / 2, by + bh / 2
        gap = 0.05
        xlo, xhi = max(ax, bx), min(ax + aw, bx + bw)
        ylo, yhi = max(ay, by), min(ay + ah, by + bh)
        if xhi - xlo > 0.2:
            mid = (xlo + xhi) / 2
            down = bcy > acy
            p1 = (mid, (ay + ah + gap) if down else (ay - gap))
            p2 = (mid, (by - gap) if down else (by + bh + gap))
        elif yhi - ylo > 0.2:
            right = bcx > acx
            mid = (ylo + yhi) / 2
            p1 = ((ax + aw + gap) if right else (ax - gap), mid)
            p2 = ((bx - gap) if right else (bx + bw + gap), mid)
        else:
            down = bcy > acy
            p1 = (acx, (ay + ah + gap) if down else (ay - gap))
            p2 = (bcx, (by - gap) if down else (by + bh + gap))
        T.connect(sl, p1[0], p1[1], p2[0], p2[1], T.ACCENT, 1.75)
        span = abs(p2[1] - p1[1]) if p1[0] == p2[0] else abs(p2[0] - p1[0])
        if len(e) > 2 and e[2] and span > 0.9 and p1[1] == p2[1]:
            tf = T.txt(sl, (p1[0] + p2[0]) / 2 - span / 2 + 0.05, p1[1] - 0.22, span - 0.1, 0.17,
                       align=PP_ALIGN.CENTER)
            T.para(tf, e[2], size=8, color=T.SLATE, space_after=0, first=True)

    for n in s["nodes"]:
        bx, by, bw, bh = box[n["id"]]
        T.card(sl, bx, by, bw, bh, n.get("num", ""), n["head"], n.get("body", ""))
    T.notes(sl, s.get("source"), s.get("talk"))


def act(s: dict[str, Any]) -> None:
    """A short act divider: the act number as a watermark, a title and one line."""
    sl = T.act_divider(s["num"], s["title"], s["sub"])
    T.notes(sl, s.get("source"), s.get("talk"))


def _centre(sl: Any, top: float, h: float, used: float) -> None:
    """Move the shape just drawn so its text sits in the middle of ``h``."""
    from pptx.util import Inches

    sl.shapes[-1].top = Inches(top + max(0.0, (h - used) / 2))


def qa(s: dict[str, Any]) -> None:
    """Questions and their short answers, one tinted row each (the executive summary)."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    rows = s["rows"]
    gap = 0.14
    rh = (bottom - y - gap * (len(rows) - 1)) / len(rows)
    qw = s.get("q_w", 3.0)
    for i, (q, a) in enumerate(rows):
        top = y + i * (rh + gap)
        T.rect(sl, T.ML, top, T.CW, rh, fill=T.WASH_L)
        T.rect(sl, T.ML, top, qw, rh, fill=T.WASH)
        T.rect(sl, T.ML, top, 0.06, rh, fill=T.ACCENT)
        used = T.fitted(sl, T.ML + 0.25, top + 0.08, qw - 0.4, rh - 0.16,
                        lambda tf, z, q=q: T.para(tf, q, size=z, color=T.ACCENT_D, bold=True,
                                                  font=HEAD, first=True, space_after=0, line=1.1),
                        s.get("q_size", 16), 10)
        _centre(sl, top + 0.08, rh - 0.16, used)
        x = T.ML + qw + 0.3
        used = T.fitted(sl, x, top + 0.08, T.ML + T.CW - x - 0.2, rh - 0.16,
                        lambda tf, z, a=a: T.para(tf, a, size=z, color=T.INK, first=True,
                                                  space_after=0, line=1.18),
                        s.get("size", 15), 9.5)
        _centre(sl, top + 0.08, rh - 0.16, used)
    T.notes(sl, s.get("source"), s.get("talk"))


def _step_size(items: list[tuple[str, str]], width: float, height: float, start: float,
               floor: float = 9.5) -> float:
    """The largest size at which every numbered item fits its slot: one size for all, so the
    list reads as one list."""
    size = start
    while size >= floor:
        need = max(text_h(lead, width * 0.94, size + 1, True, HEAD, 1.1)
                   + text_h(body, width * 0.94, size - 1, False, SANS, 1.2) + 0.08
                   for lead, body in items)
        if need <= height - 0.06:
            return size
        size -= 0.5
    raise T.DoesNotFit(f"numbered items do not fit {width:.2f}x{height:.2f}")


def steps(s: dict[str, Any]) -> None:
    """Numbered items -- principles or the steps of a procedure -- each a disc, a bold lead
    and a line or two under it, in one or two columns."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    items = s["items"]
    cols = s.get("cols", 1)
    per = (len(items) + cols - 1) // cols
    colw = (T.CW - 0.4 * (cols - 1)) / cols
    rh = (bottom - y) / per
    d = min(0.5, rh - 0.12)
    tw = colw - d - 0.25
    size = _step_size(items, tw, rh - 0.1, s.get("size", 15.5))
    for i, (lead, body) in enumerate(items):
        c, r = divmod(i, per)
        x = T.ML + c * (colw + 0.4)
        top = y + r * rh
        if r:
            T.rect(sl, x + d + 0.25, top - 0.004, tw, 0.008, fill=T.RULE)
        T.circle(sl, x, top + 0.06, d, str(i + 1), size=min(15, d * 30))

        def write(tf: Any, z: float, lead: str = lead, body: str = body) -> None:
            T.para(tf, lead, size=z + 1, color=T.INK, bold=True, font=HEAD, first=True,
                   space_after=1, line=1.1)
            T.para(tf, body, size=z - 1, color=T.SLATE, space_after=0, line=1.2)

        T.fitted(sl, x + d + 0.25, top + 0.06, tw, rh - 0.1, write, size, 9)
    T.notes(sl, s.get("source"), s.get("talk"))


def shot(s: dict[str, Any]) -> None:
    """A real screenshot, as large as the slide allows, with what to look at beside it."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    iw = T.CW * s.get("img_w", 0.66)
    px, py, pw, ph = T.picture(sl, str(T.ROOT / s["image"]), T.ML, y, iw, bottom - y)
    x = px + pw + 0.35
    _body(sl, x, y, T.ML + T.CW - x, bottom - y, s, 14)
    T.notes(sl, s.get("source"), s.get("talk"))


def shots(s: dict[str, Any]) -> None:
    """Two screenshots side by side, each with a bold caption and a line under it."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    w = (T.CW - 0.4) / 2
    cap = s.get("cap_h", 1.0)
    for i, (image, head, body) in enumerate(s["images"]):
        x = T.ML + i * (w + 0.4)
        px, py, pw, ph = T.picture(sl, str(T.ROOT / image), x, y, w, bottom - y - cap - 0.12)

        def write(tf: Any, z: float, head: str = head, body: str = body) -> None:
            T.para(tf, head, size=z + 1.5, color=T.ACCENT_D, bold=True, font=HEAD, first=True,
                   space_after=2, line=1.1)
            T.para(tf, body, size=z, color=T.SLATE, space_after=0, line=1.2)

        T.fitted(sl, x, py + ph + 0.14, w, bottom - (py + ph + 0.14), write, 13.5, 9)
    T.notes(sl, s.get("source"), s.get("talk"))


def tree(s: dict[str, Any]) -> None:
    """A flow drawn as a numbered step tree: who talks to whom in a band, the steps in a
    code panel, and what follows from it beside them."""
    sl, y = T.content(s["title"], s.get("kicker"))
    y = _intro(sl, y, s.get("intro"))
    bottom = _note(sl, s.get("note"), s.get("takeaway"))
    frac = s.get("code_w", 0.62) if s.get("items") or s.get("rows") else 1.0
    cw = T.CW * frac - (0.2 if frac < 1 else 0)
    T.rect(sl, T.ML, y, cw, 0.42, fill=T.ACCENT_D)
    tf = T.txt(sl, T.ML, y + 0.09, cw, 0.28, align=PP_ALIGN.CENTER)
    T.para(tf, s["actors"], size=13, color=T.WHITE, bold=True, first=True, space_after=0)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    top = y + 0.42
    T.code(sl, T.ML, top, cw, bottom - top, s["code"], s.get("code_size", 12.5), steps=True)
    if frac < 1:
        x = T.ML + T.CW * frac + 0.2
        _body(sl, x, y, T.ML + T.CW - x, bottom - y, s, 14)
    T.notes(sl, s.get("source"), s.get("talk"))


def thanks(s: dict[str, Any]) -> None:
    """The closing slide."""
    T._state["n"] += 1
    sl = T.blank()
    T.rect(sl, 0, 0, SW, SH, fill=T.ACCENT_D)
    T.rect(sl, 0, 0, 0.20, SH, fill=T.ACCENT_DD)
    T.mark(sl, (SW - 1.4) / 2, 1.0, 1.4, light=True)
    tf = T.txt(sl, T.ML, 2.65, T.CW, 1.0, align=PP_ALIGN.CENTER)
    T.para(tf, s["title"], size=48, color=T.WHITE, bold=True, font=HEAD, first=True,
           space_after=0, line=1.05)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    T.rect(sl, (SW - 1.7) / 2, 3.72, 1.7, 0.035, fill=T.ACCENT_L)
    tf = T.txt(sl, T.ML, 3.95, T.CW, 0.55, align=PP_ALIGN.CENTER)
    T.para(tf, s["sub"], size=24, color=T.WHITE, italic=True, font=HEAD, first=True,
           space_after=0)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf = T.txt(sl, T.ML, 4.85, T.CW, 1.3, align=PP_ALIGN.CENTER)
    for i, line in enumerate(s["lines"]):
        p = T.para(tf, line, size=18 if i == 0 else 13, color=T.WHITE if i == 0 else T.WASH,
                   bold=i == 0, first=i == 0, space_after=3)
        p.alignment = PP_ALIGN.CENTER
    T.notes(sl, s.get("source"), s.get("talk"))


KINDS = {
    "title": title,
    "divider": divider,
    "bullets": bullets,
    "table": table,
    "cards": cards,
    "stats": stats,
    "split": split,
    "flow": flow,
    "code": code,
    "context": context,
    "act": act,
    "qa": qa,
    "steps": steps,
    "shot": shot,
    "shots": shots,
    "tree": tree,
    "thanks": thanks,
}


def render(slides: list[dict[str, Any]]) -> None:
    for i, spec in enumerate(slides, 1):
        try:
            KINDS[spec["kind"]](spec)
        except T.DoesNotFit as exc:
            raise T.DoesNotFit(f"slide {i} ({spec.get('title')!r}): {exc}") from exc
