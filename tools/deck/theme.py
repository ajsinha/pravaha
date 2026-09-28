"""
The deck design system: Pravaha's verdigris (brand/README.md: accent #0E7C7B,
accent-deep #0A5C5B, accent-wash #DCEBEA, ink #0F1A1C, ground #EEF1F2) and
IBM Plex -- Sans Condensed for headings, Sans for text; Consolas for code -- with
the layout primitives every slide is drawn with.

Every primitive that places text measures it with ``metrics`` -- the same
estimator the geometry audit uses -- and the fitting helpers shrink the type or
refuse, so a slide that would overflow fails the build rather than the reader.

The alert colour is semantic only (brand/README.md), so it never decorates a
slide. A retraction is not an error: where a -1 is drawn it wears the indigo the
console gives retractions (console/web/templates/base.html, ``--retract``).

Project Pravaha -- Ask once. Answer always.
Copyright (c) 2026 Ashutosh Sinha <ajsinha@gmail.com>. All rights reserved.
PROPRIETARY AND CONFIDENTIAL. See LICENSE at the repository root.
"""

from __future__ import annotations

from typing import Any, Callable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from pptx.util import Inches as In

from metrics import FOOTER_Y, HEAD, MONO, SAFETY, SANS, SH, SW, est_lines, text_extent, text_h

ACCENT = RGBColor(0x0E, 0x7C, 0x7B)
ACCENT_D = RGBColor(0x0A, 0x5C, 0x5B)
ACCENT_DD = RGBColor(0x07, 0x44, 0x43)
ACCENT_L = RGBColor(0x37, 0xB0, 0xA8)
WASH = RGBColor(0xDC, 0xEB, 0xEA)
WASH_L = RGBColor(0xEE, 0xF6, 0xF5)
INK = RGBColor(0x0F, 0x1A, 0x1C)
SLATE = RGBColor(0x3E, 0x4D, 0x50)
MUTED = RGBColor(0x5F, 0x6E, 0x71)
RULE = RGBColor(0xCF, 0xDA, 0xDA)
GROUND = RGBColor(0xEE, 0xF1, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RETRACT = RGBColor(0x29, 0x33, 0x52)
CODE_BG = RGBColor(0x0F, 0x1A, 0x1C)
CODE_FG = RGBColor(0xDC, 0xEB, 0xEA)
CODE_KW = RGBColor(0x6F, 0xD0, 0xC8)

ML = 0.85
CW = SW - 2 * ML
BODY_BOTTOM = FOOTER_Y - 0.12

_state: dict[str, Any] = {"prs": None, "chapter": "", "n": 0}


class DoesNotFit(Exception):
    """A block that cannot be made to fit its box at the smallest permitted size."""


def new_deck(chapter: str) -> Any:
    """Start a fresh 16:9 presentation and make it the current one."""
    prs = Presentation()
    prs.slide_width = In(SW)
    prs.slide_height = In(SH)
    _state.update(prs=prs, chapter=chapter, n=0)
    return prs


def blank() -> Any:
    prs = _state["prs"]
    return prs.slides.add_slide(prs.slide_layouts[6])


def notes(sl: Any, text: str | None) -> None:
    """The speaker notes: where every figure and claim on the slide comes from."""
    if text:
        sl.notes_slide.notes_text_frame.text = text


def remove(shape: Any) -> None:
    el = shape._element
    el.getparent().remove(el)


def rect(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: Any = None,
    line: Any = None,
    lw: float = 1.0,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    s = sl.shapes.add_shape(shape, In(x), In(y), In(w), In(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    s.shadow.inherit = False
    return s


def txt(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
) -> Any:
    tb = sl.shapes.add_textbox(In(x), In(y), In(w), In(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    tf.paragraphs[0].alignment = align
    return tf


def para(
    tf: Any,
    text: str,
    size: float = 14,
    color: Any = INK,
    bold: bool = False,
    font: str = SANS,
    italic: bool = False,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    space_before: float = 0,
) -> Any:
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def runs(
    tf: Any,
    parts: list[tuple],
    size: float = 14,
    space_after: float = 6,
    first: bool = False,
    line: float = 1.2,
    level: int = 0,
    space_before: float = 0,
    font: str = SANS,
) -> Any:
    """``parts`` is a list of ``(text, color, bold)`` tuples, one run each."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line
    p.level = level
    for text, color, bold in parts:
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        r.font.name = font
    return p


def fitted(
    sl: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    write: Callable[[Any, float], None],
    start: float,
    floor: float = 9.0,
) -> float:
    """Write a text block at the largest size in [floor, start] whose estimated
    extent fits ``h``; return the height it uses. Raises ``DoesNotFit``."""
    size = start
    while size >= floor:
        tf = txt(sl, x, y, w, h)
        write(tf, size)
        need = text_extent(sl.shapes[-1])
        if need <= h:
            return need
        remove(sl.shapes[-1])
        size -= 0.5
    raise DoesNotFit(f"text block does not fit {w:.2f}x{h:.2f} at {floor}pt")


def mark(sl: Any, x: float, y: float, size: float, light: bool = False) -> None:
    """The flow mark (brand/mark.svg): three streamlines in phase, the outer two at
    55 % opacity, drawn as polylines sampled from the SVG's quadratic curves so the
    deck carries no image file. Stroke 7 units in a 64-unit frame, round caps."""
    scale = size / 64.0

    def curve(y0: float) -> list[tuple[float, float]]:
        # M4,y0 Q16,y0-12 28,y0 T52,y0 -- two quadratic segments, the second's
        # control point the reflection of the first's.
        segs = [((4, y0), (16, y0 - 12), (28, y0)), ((28, y0), (40, y0 + 12), (52, y0))]
        pts: list[tuple[float, float]] = []
        for p0, c, p1 in segs:
            for i in range(0 if not pts else 1, 17):
                t = i / 16
                px = (1 - t) ** 2 * p0[0] + 2 * (1 - t) * t * c[0] + t**2 * p1[0]
                py = (1 - t) ** 2 * p0[1] + 2 * (1 - t) * t * c[1] + t**2 * p1[1]
                pts.append((x + px * scale, y + py * scale))
        return pts

    colour = ACCENT_L if light else ACCENT
    for y0, alpha in ((18, 55), (32, 100), (46, 55)):
        pts = curve(y0)
        fb = sl.shapes.build_freeform(In(pts[0][0]), In(pts[0][1]), scale=1.0)
        fb.add_line_segments([(In(px), In(py)) for px, py in pts[1:]], close=False)
        shp = fb.convert_to_shape()
        shp.fill.background()
        shp.line.color.rgb = colour
        shp.line.width = Pt(7 * scale * 72)
        shp.shadow.inherit = False
        ln = shp.line._get_or_add_ln()
        ln.set("cap", "rnd")
        if alpha < 100:
            clr = ln.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = clr.makeelement(qn("a:alpha"), {"val": str(alpha * 1000)})
            clr.append(a)


def footer(sl: Any) -> None:
    rect(sl, ML, FOOTER_Y, CW, 0.008, fill=RULE)
    tf = txt(sl, ML, SH - 0.46, CW * 0.7, 0.24)
    para(tf, _state["chapter"], size=8.5, color=MUTED, first=True, space_after=0)
    tf = txt(sl, ML + CW * 0.7, SH - 0.46, CW * 0.3, 0.24, align=PP_ALIGN.RIGHT)
    para(tf, str(_state["n"]), size=8.5, color=MUTED, first=True, space_after=0, bold=True)


def content(title: str, kicker: str | None = None) -> tuple[Any, float]:
    """A content slide's chrome. Returns ``(slide, body_top)``."""
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=WHITE)
    rect(sl, 0, 0.62, 0.30, 0.055, fill=ACCENT)
    y = 0.52
    if kicker:
        tf = txt(sl, ML, y, CW, 0.24)
        para(tf, kicker.upper(), size=10.5, color=ACCENT, bold=True, first=True, space_after=0)
        y += 0.30
    size = 27.0
    while size > 20 and est_lines(title, CW * SAFETY, size, True, HEAD) > 1:
        size -= 1
    th = text_h(title, CW * SAFETY, size, True, HEAD, 1.15)
    tf = txt(sl, ML, y, CW, th + 0.04)
    para(tf, title, size=size, color=INK, font=HEAD, bold=True, first=True, space_after=0, line=1.15)
    body_top = y + th + 0.24
    rect(sl, ML, body_top - 0.14, CW, 0.012, fill=RULE)
    footer(sl)
    return sl, body_top


def divider(num: str, title: str, sub: str, points: list[str]) -> Any:
    """A verdigris part divider with its contents on the right."""
    _state["chapter"] = f"Part {num} · {title}"
    _state["n"] += 1
    sl = blank()
    rect(sl, 0, 0, SW, SH, fill=ACCENT_D)
    rect(sl, 0, 0, 0.18, SH, fill=ACCENT_DD)
    mark(sl, ML + 0.25, 0.85, 0.9, light=True)
    tf = txt(sl, ML + 0.25, 2.0, CW * 0.58, 0.4)
    para(tf, f"PART {num}", size=12, color=ACCENT_L, bold=True, first=True, space_after=0)
    tw = CW * 0.58
    size = 40.0
    while size > 24 and est_lines(title, tw * SAFETY, size, True, HEAD) > 2:
        size -= 2
    th = text_h(title, tw * SAFETY, size, True, HEAD, 1.05)
    tf = txt(sl, ML + 0.25, 2.45, tw, th + 0.05)
    para(tf, title, size=size, color=WHITE, font=HEAD, bold=True, first=True, space_after=0,
         line=1.05)
    ry = 2.45 + th + 0.25
    rect(sl, ML + 0.25, ry, 1.5, 0.035, fill=ACCENT_L)
    fitted(
        sl,
        ML + 0.25,
        ry + 0.25,
        tw,
        6.3 - (ry + 0.25),
        lambda tf, s: para(
            tf, sub, size=s, color=WASH, italic=True, first=True, space_after=0, line=1.3
        ),
        15,
        10.5,
    )
    x = ML + CW * 0.66
    tf = txt(sl, x, 2.0, CW * 0.34, 0.3)
    para(tf, "IN THIS PART", size=9.5, color=ACCENT_L, bold=True, first=True, space_after=0)

    def write(tf: Any, s: float) -> None:
        for i, pnt in enumerate(points):
            para(tf, pnt, size=s, color=WASH, space_after=6, line=1.15, first=i == 0)

    fitted(sl, x, 2.4, CW * 0.34, 4.0, write, 12.5, 9)
    return sl


def _row_heights(
    data: list[list[str]], widths: list[float], fs: float, hfs: float, bold_col0: bool
) -> list[float]:
    heights = []
    for r, row in enumerate(data):
        size = hfs if r == 0 else fs
        need = 0.0
        for c, cell in enumerate(row):
            bold = r == 0 or (bold_col0 and c == 0)
            need = max(need, text_h(cell, (widths[c] - 0.18) * SAFETY, size, bold, SANS, 1.0))
        heights.append(max(0.30, need + 0.12))
    return heights


def _fill_cell(
    cell: Any, text: str, r: int, c: int, fs: float, hfs: float, bold_col0: bool
) -> None:
    cell.margin_left = In(0.09)
    cell.margin_right = In(0.07)
    cell.margin_top = cell.margin_bottom = In(0.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_D if r == 0 else (WASH_L if r % 2 == 0 else WHITE)
    tf = cell.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = str(text)
    run.font.name = SANS
    run.font.size = Pt(hfs if r == 0 else fs)
    run.font.bold = r == 0 or (bold_col0 and c == 0)
    run.font.color.rgb = WHITE if r == 0 else (ACCENT_D if bold_col0 and c == 0 else INK)


def table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    col_w: list[float] | None = None,
    fs: float = 11.5,
    hfs: float = 11,
    bold_col0: bool = True,
) -> float:
    """A table whose row heights are computed from its wrapped text.
    Returns the rendered height, so callers place what follows from it."""
    cols = len(data[0])
    col_w = col_w or [1.0] * cols
    widths = [w * c / sum(col_w) for c in col_w]
    heights = _row_heights(data, widths, fs, hfs, bold_col0)
    total = sum(heights)
    gf = sl.shapes.add_table(len(data), cols, In(x), In(y), In(w), In(total))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(widths):
        tbl.columns[i].width = Emu(int(In(cw)))
    for r, row in enumerate(data):
        tbl.rows[r].height = Emu(int(In(heights[r])))
        for c, cell in enumerate(row):
            _fill_cell(tbl.cell(r, c), cell, r, c, fs, hfs, bold_col0)
    return total


def fitted_table(
    sl: Any,
    data: list[list[str]],
    x: float,
    y: float,
    w: float,
    h: float,
    col_w: list[float] | None = None,
    start: float = 14.0,
    bold_col0: bool = True,
) -> float:
    """``table`` at the largest size in [9, start] that fits ``h``."""
    fs = start
    while fs >= 9.0:
        used = table(sl, data, x, y, w, col_w, fs=fs, hfs=min(fs, 13), bold_col0=bold_col0)
        if used <= h:
            return used
        remove(sl.shapes[-1])
        fs -= 0.5
    raise DoesNotFit(f"table of {len(data)} rows does not fit {h:.2f}in")


def card(sl: Any, x: float, y: float, w: float, h: float, num: str, title: str, body: str) -> None:
    """A bordered card whose contents are fitted inside the border."""
    rect(sl, x, y, w, h, fill=WHITE, line=RULE, lw=0.9)
    rect(sl, x, y, w, 0.055, fill=ACCENT)
    inner = w - 0.40
    top = y + 0.18
    if num:
        tf = txt(sl, x + 0.20, top, inner, 0.24)
        para(tf, num, size=10, color=ACCENT, bold=True, first=True, space_after=0)
        top += 0.28
    used = fitted(
        sl,
        x + 0.20,
        top,
        inner,
        min(0.78, y + h - 0.2 - top),
        lambda tf, s: para(
            tf,
            title,
            size=s,
            color=INK,
            bold=True,
            font=HEAD,
            first=True,
            space_after=0,
            line=1.1,
        ),
        16.5,
        10.5,
    )
    by = top + used + 0.10
    if body:
        fitted(
            sl,
            x + 0.20,
            by,
            inner,
            (y + h - 0.14) - by,
            lambda tf, s: para(tf, body, size=s, color=SLATE, first=True, space_after=0, line=1.2),
            14,
            8.5,
        )


def statbar(sl: Any, y: float, stats: list[tuple[str, str]], h: float = 1.25) -> None:
    gap = 0.22
    bw = (CW - gap * (len(stats) - 1)) / len(stats)
    for i, (big, label) in enumerate(stats):
        x = ML + i * (bw + gap)
        rect(sl, x, y, bw, h, fill=WASH_L)
        rect(sl, x, y, 0.045, h, fill=ACCENT)
        fitted(
            sl,
            x + 0.22,
            y + 0.14,
            bw - 0.34,
            0.5,
            lambda tf, s, b=big: para(
                tf,
                b,
                size=s,
                color=ACCENT_D,
                bold=True,
                font=HEAD,
                first=True,
                space_after=0,
                line=1.0,
            ),
            25,
            13,
        )
        fitted(
            sl,
            x + 0.22,
            y + 0.66,
            bw - 0.34,
            h - 0.74,
            lambda tf, s, lab=label: para(
                tf, lab, size=s, color=SLATE, first=True, space_after=0, line=1.12
            ),
            10.5,
            8,
        )


def code(sl: Any, x: float, y: float, w: float, h: float, lines: list[str], start: float = 13.0,
         floor: float = 8.5) -> float:
    """A dark code panel. Every line is its own paragraph and none may wrap: the
    size is the largest at which the longest line fits the width and all the
    lines fit the height. Returns the panel's height."""
    pad = 0.22
    inner_w = w - 2 * pad
    size = start
    while size >= floor:
        fits_w = all(est_lines(ln or " ", inner_w * SAFETY, size, False, MONO) == 1 for ln in lines)
        need = sum(text_h(ln or " ", inner_w * SAFETY, size, False, MONO, 1.15) for ln in lines)
        need += 2 * pad
        if fits_w and need <= h:
            break
        size -= 0.5
    else:
        raise DoesNotFit(f"code block of {len(lines)} lines does not fit {w:.2f}x{h:.2f}")
    rect(sl, x, y, w, need, fill=CODE_BG)
    rect(sl, x, y, 0.05, need, fill=ACCENT_L)
    tf = txt(sl, x + pad, y + pad, inner_w, need - 2 * pad)
    for i, ln in enumerate(lines):
        colour = CODE_KW if ln.lstrip().startswith(("--", "#", "//")) else CODE_FG
        para(tf, ln if ln else " ", size=size, color=colour, font=MONO, first=i == 0,
             space_after=0, line=1.0)
    return need


def connect(
    sl: Any, x1: float, y1: float, x2: float, y2: float, color: Any = SLATE, width: float = 1.5
) -> Any:
    """A straight connector with an arrowhead. Elbow connectors auto-route into
    detours that collide with the nodes they join; do not use them."""
    c = sl.shapes.add_connector(1, In(x1), In(y1), In(x2), In(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c
